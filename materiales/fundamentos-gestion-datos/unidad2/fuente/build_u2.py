"""Genera: Unidad 2 - Python para Datos, Limpieza y ML Básico (Fundamentos de Gestión de Datos, TECSUP 2026-II)."""
import json, os, sys
import pandas as pd
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from fmt import Doc, add_runs, NAVY, TEAL, GRAY, TEXT_W, _shade, _borders, _cell_margins
from runner import run_all
from snippets import S

PAGES = json.load(open('pages.json')) if os.path.exists('pages.json') else {}
OUT = sys.argv[1] if len(sys.argv) > 1 else 'Unidad2.docx'
ns, OUTS = run_all()

# --- Verificación: las cifras citadas en el texto coinciden con los resultados reales ---
txt = lambda k: OUTS[k][0]
assert '(408, 8)' in txt('s5_conexion') and '(3004, 6)' in txt('s5_join')
assert 'Clientes con email repetido: 8' in txt('s6_dup_cli') and 'Pedidos duplicados: 15' in txt('s6_dup_ped')
assert 'Atípicos por IQR: 257 | por z-score (|z|>3): 15' in txt('s6_iqr')
assert 'Coeficiente visitas_web_mes: 233.74' in txt('s7_simple') and 'R² (prueba): 0.462' in txt('s7_simple')
assert 'RMSE: 1447.2' in txt('s7_simple') and 'R² (prueba): 0.49' in txt('s7_multiple')
assert 'Exactitud (accuracy): 0.825' in txt('s7_clasif')
assert 'Iteraciones hasta converger: 7' in txt('s8_kmeans')
g = OUTS['s8_guardar'][1]
assert round(g.gasto_total.sum(), 1) == 840813.6 and g.iloc[0].clientes == 171

D = Doc()
d = D.d


def fmtv(v):
    if isinstance(v, float):
        return f'{v:,.2f}'
    return str(v)


def render(val, index=True, size=8.2, widths=None, index_name=None, split=False):
    if split and isinstance(val, pd.Series):
        half = (len(val) + 1) // 2
        items = [('NaN' if pd.isna(k) else repr(k).strip("'") if isinstance(k, str) and k != k.strip() else str(k), str(v))
                 for k, v in val.items()]
        left, right = items[:half], items[half:] + [('', '')] * (2 * half - len(items))
        rows = [[a[0], a[1], b_[0], b_[1]] for a, b_ in zip(left, right)]
        name = index_name or 'valor'
        D.table([name, 'count', name, 'count'], rows, [3.4, 1.8, 3.4, 1.8], size=size, align=['l', 'r', 'l', 'r'])
        return
    if isinstance(val, pd.Series):
        val = val.to_frame(name=val.name if val.name not in (None, 0) else 'valor')
        index = True
    df = val.reset_index() if index else val
    if index and index_name:
        df = df.rename(columns={df.columns[0]: index_name})
    cols = [str(c) for c in df.columns]
    rows = [[fmtv(v) if not (isinstance(v, float) and pd.isna(v)) else 'NaN' for v in r] for r in df.itertuples(index=False)]
    if widths is None:
        lens = [max(len(c), *(len(x[i]) for x in rows)) + 2 for i, c in enumerate(cols)]
        tot = sum(lens); full = min(TEXT_W, 1.2 + 0.19 * tot)
        widths = [full * l / tot for l in lens]
    align = ['r' if pd.api.types.is_numeric_dtype(df[c]) else 'l' for c in df.columns]
    D.table(cols, rows, widths, size=size, align=align)


def salida(text):
    D.code(text.rstrip('\n'), size=8.0, fill='FFFFFF', edge='BFBFBF')


def py_case(pregunta, keys, interpretacion, index=True, widths=None, index_name=None, show=True, split=False):
    keys = [keys] if isinstance(keys, str) else keys
    if pregunta:
        D.label('Pregunta de negocio')
        D.p(f'**{pregunta}**', color=NAVY, after=3, keep=True)
    D.label('Código Python')
    D.code('\n\n'.join(S[k] for k in keys), size=8.2)
    if show:
        D.label('Resultado')
        for k in keys:
            t, v = OUTS[k]
            if t.strip():
                salida(t)
            if v is not None:
                render(v, index=index, widths=widths, index_name=index_name, split=split)
    if interpretacion:
        D.label('Interpretación')
        for t in ([interpretacion] if isinstance(interpretacion, str) else interpretacion):
            D.p(t, align='j')


def cierre(semana, dominar, preguntas, conexion=None, siguiente=None):
    h = D.h2(f'Cierre de la Semana {semana}')
    h.paragraph_format.space_before = Pt(6)
    D.box('Lo que debes dominar', bullets=dominar, size=9, fill='EAF4EA', edge='38761D', title_color=RGBColor(0x27, 0x5D, 0x14))
    D.box('Comprueba tu aprendizaje', size=9, fill='FFF7E6', edge='C58B00', title_color=RGBColor(0x8A, 0x5A, 0x00),
          lines=[f'**{k}.** {q}' for k, q in enumerate(preguntas, 1)])
    if conexion:
        D.box(siguiente or f'Conexión con la Semana {semana + 1}', lines=conexion, fill='EEF3FA', size=9)
    last = d.paragraphs[-1]._p
    last.getparent().remove(last)


# =====================================================================
# PORTADA
# =====================================================================
for _ in range(3):
    D.spacer(12)
t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
_cell_margins(t, 400, 400, 500, 500)
c = t.rows[0].cells[0]; c.width = Cm(TEXT_W)
t._tbl.tblGrid.findall(qn('w:gridCol'))[0].set(qn('w:w'), str(int(TEXT_W * 567)))
_shade(c._tc.get_or_add_tcPr(), '1F3864'); _borders(c._tc.get_or_add_tcPr(), {'left': (48, '2E75B6')})
for k, (tx, sz, col, b, after) in enumerate([
        ('TECSUP · Carrera de Big Data y Ciencia de Datos', 10.5, 'BDD7EE', False, 18),
        ('FUNDAMENTOS DE GESTIÓN DE DATOS', 22, 'FFFFFF', True, 10),
        ('UNIDAD 2', 14, '9DC3E6', True, 2),
        ('PYTHON PARA DATOS, LIMPIEZA Y MACHINE LEARNING BÁSICO', 16, 'FFFFFF', True, 16),
        ('Material de estudio · Semanas 5 a 8', 11, 'DEEAF6', False, 2),
        ('De la base de datos a la decisión basada en modelos', 11, 'DEEAF6', False, 0)]):
    para = c.paragraphs[0] if k == 0 else c.add_paragraph()
    r = para.add_run(tx); r.font.size = Pt(sz); r.bold = b; r.font.color.rgb = RGBColor.from_string(col)
    para.paragraph_format.space_after = Pt(after)
D.spacer(30)
D.table(['Semana', 'Tema'], [['5', 'Python y Pandas para Datos'], ['6', 'Limpieza de Datos'],
                             ['7', 'ML Supervisado: Regresión y Clasificación'], ['8', 'ML No Supervisado: Clustering (K-Means)']],
        [2.5, 10.5], size=10, align=['c', 'l'])
D.spacer(40)
for tx in ('**Docente:** Pilar Rocío Sayán Mejía', '**Periodo académico:** 2026-II'):
    D.p(tx, size=11, align='c', after=3)

# =====================================================================
# CONTENIDO
# =====================================================================
TOC = ['Introducción a la Unidad 2', 'Semana 5 · Python y Pandas para Datos', 'Semana 6 · Limpieza de Datos',
       'Semana 7 · ML Supervisado: Regresión y Clasificación', 'Semana 8 · ML No Supervisado: Clustering con K-Means',
       'Integración de la Unidad 2', 'Conclusiones', 'Referencias bibliográficas']
TOC_SUB = {
    'Introducción a la Unidad 2': ['¿Qué aprenderemos?', 'Caso transversal: RutaMarket, un año después'],
    'Semana 5 · Python y Pandas para Datos': ['5.1 Google Colab: el entorno de trabajo', '5.2 Lo esencial de Python para datos',
        '5.3 pandas: DataFrame y Series', '5.4 Leer datos desde SQLite y CSV', '5.5 Explorar los datos: el análisis exploratorio (EDA)',
        '5.6 Visualizar: matplotlib y seaborn'],
    'Semana 6 · Limpieza de Datos': ['6.1 ¿Por qué limpiar?', '6.2 Diagnóstico inicial', '6.3 Tipos incorrectos y valores imposibles',
        '6.4 Inconsistencias de texto', '6.5 Registros duplicados', '6.6 Datos faltantes: detección e imputación',
        '6.7 Valores atípicos: boxplot, z-score e IQR', '6.8 Normalización: Min-Max y estandarización', '6.9 Guardar el dataset limpio en SQLite'],
    'Semana 7 · ML Supervisado: Regresión y Clasificación': ['7.1 ¿Qué es Machine Learning?', '7.2 El proceso CRISP-DM',
        '7.3 Datos de entrenamiento y de prueba', '7.4 Regresión lineal simple', '7.5 Regresión lineal múltiple',
        '7.6 Overfitting y underfitting', '7.7 Clasificación: concepto e interpretación', '7.8 Uso responsable de los modelos'],
    'Semana 8 · ML No Supervisado: Clustering con K-Means': ['8.1 Aprendizaje no supervisado: ¿cuándo usar clustering?',
        '8.2 Cómo funciona K-Means', '8.3 Preparar los datos y elegir k: el método del codo', '8.4 Interpretar los clusters',
        '8.5 Guardar los segmentos en la base de datos', '8.6 Limitaciones de K-Means'],
    'Integración de la Unidad 2': ['Mapa del recorrido', 'Caso integrador', 'Relación con el proyecto PMD1', 'Checklist de dominio de la Unidad 2'],
}
D.h1('Contenido')
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
for title in TOC:
    for lvl, tt in [(0, title)] + [(1, s) for s in TOC_SUB.get(title, [])]:
        para = d.add_paragraph(); pf = para.paragraph_format
        pf.space_after = Pt(1.5 if lvl else 1); pf.space_before = Pt(0 if lvl else 5); pf.left_indent = Cm(0.7 * lvl)
        pf.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = para.add_run(tt); r.font.size = Pt(10 if lvl == 0 else 9.3); r.bold = lvl == 0
        if lvl == 0:
            r.font.color.rgb = NAVY
        r2 = para.add_run('\t' + str(PAGES.get(tt, '–'))); r2.font.size = r.font.size; r2.bold = lvl == 0

# =====================================================================
# INTRODUCCIÓN
# =====================================================================
D.h1('Introducción a la Unidad 2')
D.p('En la Unidad 1 aprendimos a comprender, organizar y consultar datos: vimos por qué la calidad importa, escribimos consultas SQL y '
    'diseñamos un modelo relacional. Esa base responde preguntas como *¿cuánto vendimos?* o *¿quiénes compraron?* La Unidad 2 da el '
    'siguiente paso: **preparar y analizar** los datos para responder preguntas más exigentes, como *¿qué problemas tienen nuestros '
    'datos?*, *¿qué variables se relacionan con el gasto de un cliente?* o *¿qué grupos de clientes existen?*', align='j')
D.p('Para ello incorporamos **Python** y su librería **pandas**, que trabajan junto a SQL: SQL extrae los datos de la base, pandas los '
    'explora y limpia, y **scikit-learn** construye modelos de *Machine Learning* básicos. El resultado vuelve a guardarse en la base de '
    'datos, cerrando el ciclo de vida del dato estudiado en la Semana 1.', align='j')
steps = [('SQLITE', 'U1'), ('PANDAS\n+ EDA', 'S5'), ('LIMPIEZA', 'S6'), ('MODELO\nSUPERVISADO', 'S7'), ('SEGMEN-\nTACIÓN', 'S8')]
t = d.add_table(rows=2, cols=9); t.alignment = WD_TABLE_ALIGNMENT.CENTER
ws = [2.9, 0.6, 2.9, 0.6, 2.9, 0.6, 2.9, 0.6, 2.9]
for i, w in enumerate(ws):
    for cell in t.columns[i].cells:
        cell.width = Cm(w)
for i, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
    gc.set(qn('w:w'), str(int(ws[i] * 567)))
for k in range(9):
    top, bot = t.rows[0].cells[k], t.rows[1].cells[k]
    for cell in (top, bot):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    if k % 2 == 0:
        name, wk = steps[k // 2]
        _shade(top._tc.get_or_add_tcPr(), ['DEEAF6', 'BDD7EE', '9DC3E6', '5B9BD5', '1F3864'][k // 2])
        for li, ln in enumerate(name.split('\n')):
            para = top.paragraphs[0] if li == 0 else top.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.space_after = Pt(0)
            r = para.add_run(ln); r.bold = True; r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if k >= 6 else NAVY
        r = bot.paragraphs[0].add_run(wk); r.font.size = Pt(8); r.font.color.rgb = GRAY
    else:
        r = top.paragraphs[0].add_run('→'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = TEAL
    top.vertical_alignment = 1
D.spacer(6)

D.h2('¿Qué aprenderemos?')
D.table(['Semana', 'Pregunta que responde', 'Al terminar, podrás…'],
        [['5. Python y Pandas', '¿Cómo llevo los datos de la base a un entorno de análisis y los exploro?',
          'Leer datos desde SQLite y CSV, explorarlos con pandas e interpretar gráficos de distribución, correlación y nulos.'],
         ['6. Limpieza de Datos', '¿Qué problemas tienen los datos y cómo los corrijo sin distorsionarlos?',
          'Detectar y tratar nulos, duplicados, tipos incorrectos y atípicos; normalizar variables y guardar el resultado en SQLite.'],
         ['7. ML Supervisado', '¿Puedo estimar un valor o una categoría a partir de otras variables?',
          'Entrenar y evaluar una regresión lineal, interpretar sus métricas y reconocer el sobreajuste y la lógica de la clasificación.'],
         ['8. ML No Supervisado', '¿Qué grupos naturales existen en mis datos?',
          'Aplicar K-Means, elegir k con el método del codo e interpretar los segmentos para decisiones de negocio.']],
        [3.4, 5.6, 8.0], size=8.7, first_col_bold=True)
D.p('**Cómo leer este material.** Como en la Unidad 1, los conceptos combinan concepto, contexto, ejemplo, aplicación, interpretación y '
    'conexión, y el código sigue el orden *pregunta de negocio → código → resultado → interpretación*. No se trata de memorizar sintaxis, '
    'sino de entender qué hace cada paso y qué significa su resultado.', align='j')

D.h2('Caso transversal: RutaMarket, un año después')
D.p('Continuamos con **RutaMarket S.A.C.**, la tienda en línea *ficticia* de la Unidad 1. Un año después, la empresa ha crecido y su base '
    'de datos `rutamarket.db` registra la operación de **2025**: clientes, productos, categorías, pedidos y detalle de pedidos, con el '
    'mismo modelo relacional diseñado en la Semana 4. Como parte de la tabla de clientes proviene de una migración desde hojas de cálculo, '
    'los datos traen problemas de calidad que deberemos descubrir y corregir. Los datos son **sintéticos**, generados con fines '
    'académicos. Todo el código fue ejecutado sobre `rutamarket.db` (pandas 2.x y scikit-learn): los resultados mostrados son los que '
    'obtendrás al reproducirlo. Los montos están en soles (S/).', align='j')
D.p('La base contiene **408 clientes** (con ciudad, edad, canal y fecha de registro, correo y visitas mensuales a la web), **3 004 pedidos** '
    '(fecha, método de pago y estado: Entregado o Devuelto), **4 737 líneas de detalle**, **20 productos** y **4 categorías**.', align='j')
D.table(['Semana', 'Cómo interviene el caso'],
        [['5', 'Leemos las tablas desde SQLite con pandas y hacemos un primer análisis exploratorio.'],
         ['6', 'Corregimos los problemas de calidad y construimos una tabla limpia de clientes (`clientes_modelo`) en la misma base.'],
         ['7', 'Estimamos el gasto anual de un cliente a partir de su actividad en la web y discutimos la clasificación de clientes de alto valor.'],
         ['8', 'Segmentamos a los clientes con K-Means y guardamos el segmento de cada uno en la base de datos.']],
        [2.0, 15.0], size=8.7, align=['c', 'l'], keep=False)


# =====================================================================
# SEMANA 5
# =====================================================================
D.h1('Semana 5 · Python y Pandas para Datos',
     subtitle='Pregunta guía: ¿cómo llevo los datos de una base de datos a un entorno de análisis y obtengo una primera visión de ellos?')
D.p('SQL es excelente para extraer y resumir datos, pero el análisis suele requerir pasos adicionales: revisar la calidad de cada columna, '
    'transformar valores, graficar y, más adelante, construir modelos. Python, con pandas, permite hacer todo eso en un mismo lugar y '
    'dejar registrado cada paso, de modo que el análisis pueda repetirse y revisarse.', align='j')

D.h2('5.1 Google Colab: el entorno de trabajo')
D.lp('Concepto', '**Google Colab** es un entorno gratuito que ejecuta código Python en servidores de Google desde el navegador, sin '
     'instalar nada en la computadora. Trabaja con **notebooks** (archivos `.ipynb`), documentos que combinan celdas de código, sus '
     'resultados y celdas de texto.')
D.table(['Tipo de celda', 'Para qué sirve', 'Ejemplo'],
        [['Código', 'Escribir y ejecutar Python (Shift + Enter)', '`clientes.head()`'],
         ['Texto (Markdown)', 'Explicar el objetivo, las decisiones y las conclusiones', '`## Hallazgos del EDA` crea un subtítulo; `**texto**`, negrita']],
        [3.2, 6.8, 7.0], size=8.7, first_col_bold=True)
D.lp('Aplicación', 'Un notebook bien hecho se lee como un informe: cada bloque de código va precedido de una celda de texto que explica '
     'qué pregunta responde y seguido de su interpretación. Las celdas se ejecutan en orden y comparten variables; si se reinicia el '
     'entorno, hay que volver a ejecutarlas desde el inicio. Para trabajar con el caso, sube `rutamarket.db` al panel *Archivos* de Colab '
     '(los archivos subidos se borran al cerrar la sesión) o léelo desde Google Drive.')

D.h2('5.2 Lo esencial de Python para datos')
D.p('Para esta unidad basta con un conjunto pequeño de elementos de Python:', align='j', keep=True)
D.code("""ciudad = 'Lima'                          # str: texto
visitas = 12                             # int: número entero
ticket = 254.80                          # float: número decimal
activo = True                            # bool: verdadero o falso
ciudades = ['Lima', 'Cusco', 'Piura']    # lista: colección ordenada
cliente = {'id': 1, 'ciudad': 'Lima'}    # diccionario: pares clave-valor
print(len(ciudades), cliente['ciudad'])  # funciones y acceso por clave""", size=8.2)
salida('3 Lima')
D.p('También usaremos **librerías**: código ya escrito que se incorpora con `import`. Las principales son **pandas** (tablas), '
    '**matplotlib** y **seaborn** (gráficos) y **scikit-learn** (modelos). En Colab ya vienen instaladas; por convención se importan con '
    'alias: `import pandas as pd`. Un **método** es una función asociada a un objeto y se invoca con un punto: `clientes.head()`.', align='j')

D.h2('5.3 pandas: DataFrame y Series')
D.lp('Concepto', 'pandas organiza los datos en dos estructuras. Un **DataFrame** es una tabla en memoria, con filas y columnas con nombre, '
     'equivalente a la tabla de una base de datos o al resultado de una consulta SQL. Una **Series** es una sola columna de un DataFrame.')
D.lp('Conexión', 'Muchas operaciones de pandas tienen un equivalente directo en lo aprendido en la Unidad 1:')
D.table(['Operación', 'SQL (Unidad 1)', 'pandas (Unidad 2)'],
        [['Ver algunas filas', 'SELECT * FROM clientes LIMIT 5;', "clientes.head()"],
         ['Filtrar filas', "WHERE ciudad = 'Lima'", "clientes[clientes['ciudad'] == 'Lima']"],
         ['Valores distintos', 'SELECT DISTINCT ciudad', "clientes['ciudad'].unique()"],
         ['Contar por grupo', 'GROUP BY ciudad + COUNT(*)', "clientes['ciudad'].value_counts()"],
         ['Agregar por grupo', 'GROUP BY id_cliente + SUM(monto)', "pedidos.groupby('id_cliente')['monto'].sum()"],
         ['Unir tablas', 'JOIN … ON …', "pedidos.merge(clientes, on='id_cliente')"]],
        [3.4, 6.0, 7.6], size=8.4, first_col_bold=True)
D.lp('Interpretación', 'SQL y pandas no compiten: SQL trabaja dentro de la base de datos y es ideal para extraer y combinar; pandas trabaja '
     'con los datos ya cargados en memoria y es ideal para explorar, transformar y preparar datos para modelos.')

D.h2('5.4 Leer datos desde SQLite y CSV')
D.p('La librería `sqlite3` abre la conexión con la base de datos y `pd.read_sql_query()` ejecuta una consulta SQL y devuelve el resultado como DataFrame.', align='j')
py_case('¿Cuántos clientes y columnas tiene la tabla de clientes, y qué aspecto tienen sus registros?', ['s5_conexion', 's5_head'],
        '`shape` indica 408 filas y 8 columnas. La vista preliminar muestra, por ejemplo, que el cliente 1 es de Lima, se registró por la web y '
        'visita la tienda 17 veces al mes. Es la misma tabla `clientes` de la Unidad 1, ahora cargada como DataFrame.', index=False)
py_case('¿Cuál es el monto total de cada pedido?', 's5_join',
        ['La consulta SQL, idéntica a las de la Semana 3, une `pedidos` con `detalle_pedido` y calcula el monto de cada pedido. El resultado '
         'llega a Python como un DataFrame de 3 004 pedidos y 6 columnas. Es la forma habitual de trabajar: **SQL prepara la tabla de '
         'análisis y pandas continúa desde ahí**.'])
py_case(None, 's5_csv', 'Los CSV se leen con `pd.read_csv()`; conviene revisar el separador (`sep`), la codificación (`encoding`) y los tipos de cada columna.', index=False)

D.h2('5.5 Explorar los datos: el análisis exploratorio (EDA)')
D.lp('Concepto', 'El **análisis exploratorio de datos** (EDA, por *Exploratory Data Analysis*) consiste en examinar un conjunto de datos '
     'para entender su estructura, detectar problemas de calidad y descubrir patrones antes de sacar conclusiones o construir modelos '
     '(Tukey, 1977). Responde preguntas como: ¿qué tipo tiene cada columna?, ¿hay vacíos?, ¿qué valores son frecuentes?, ¿hay valores extraños?')
py_case('¿Qué tipo de dato tiene cada columna y cuántos valores no nulos hay?', 's5_info',
        ['`ciudad` tiene 396 valores no nulos de 408: faltan 12. Más importante aún, `edad` aparece como `object` (texto) en lugar de número, '
         'y lo mismo ocurre con `fecha_registro`. Esto indica que algunos valores de esas columnas no son números o fechas válidas: '
         'un problema que corregiremos en la Semana 6.',
         '*Nota:* en pandas 3.x el texto se muestra con el tipo `str` en lugar de `object`; el significado es el mismo.'])
py_case('¿Cómo se distribuye el monto de los pedidos?', 's5_describe',
        ['El pedido mediano es de S/ 199.00 (`50%`) y la mitad central de los pedidos está entre S/ 99 y S/ 339. Sin embargo, el máximo es '
         'S/ 19 070 y el promedio (S/ 312.57) está muy por encima de la mediana. Un máximo tan alejado puede ser una compra real muy grande o '
         'un error de registro; el EDA no lo decide, pero **lo señala** para investigarlo.'], index_name='estadístico')
py_case('¿Cómo están registradas las ciudades?', 's5_vc',
        '`value_counts()` revela problemas de **consistencia** (Semana 1): "Lima", "lima" y " LIMA " se cuentan como ciudades distintas, '
        'y hay 12 clientes sin ciudad (`NaN`). Si calculáramos ventas por ciudad con estos datos, Lima quedaría subestimada. '
        '(Las comillas marcan los valores con espacios sobrantes.)', index_name='ciudad', split=True)

D.h2('5.6 Visualizar: matplotlib y seaborn')
D.p('Un gráfico muestra de un vistazo lo que una tabla esconde. **matplotlib** es la librería base de gráficos en Python y **seaborn** '
    'se apoya en ella para producir gráficos estadísticos con menos código. Primero resumimos los pedidos por cliente para estudiar '
    'la relación entre visitas, número de pedidos y gasto:', align='j', keep=True)
py_case(None, 's5_resumen', None, index_name='')
D.code("""import matplotlib.pyplot as plt
import seaborn as sns

fig, axs = plt.subplots(1, 3, figsize=(15, 4))
sns.histplot(pedidos['monto'], bins=40, ax=axs[0])                      # distribución
sns.heatmap(resumen[['visitas_web_mes', 'n_pedidos', 'gasto']].corr(),
            annot=True, cmap='Blues', ax=axs[1])                          # correlación
nulos = pd.concat([clientes.isnull().sum(), pedidos.isnull().sum()])
nulos.plot.barh(ax=axs[2])                                                # nulos
plt.tight_layout()""", size=8.0)
D.image('fig_eda.png', 17.0, 'Figura 1. Tres gráficos básicos del EDA de RutaMarket: distribución, correlación y valores nulos.')
D.table(['Gráfico', 'Qué muestra', 'Interpretación para RutaMarket'],
        [['Distribución (histograma)', 'Cuántos pedidos hay en cada rango de monto',
          'El 84 % de los pedidos es menor a S/ 500 y la distribución tiene una cola larga a la derecha: pocos pedidos muy grandes elevan el promedio.'],
         ['Correlación (mapa de calor)', 'Fuerza de la relación lineal entre pares de variables, de −1 a 1',
          'Los clientes que más visitan la web hacen más pedidos (0.76) y gastan más (0.55). Correlación no implica causalidad, pero orienta el modelo de la Semana 7.'],
         ['Nulos (barras)', 'Cuántos valores faltan en cada columna',
          'Faltan 90 métodos de pago y 12 ciudades. `edad` muestra 0 nulos, aunque `info()` indicó que es texto: los vacíos escritos como texto no se detectan como nulos.']],
        [3.5, 4.6, 8.9], size=8.3, first_col_bold=True)

cierre(5,
       ['Trabajar en un notebook de Colab combinando celdas de código y de texto que documentan el análisis.',
        'Usar los elementos básicos de Python (tipos, listas, diccionarios, funciones, librerías) necesarios para analizar datos.',
        'Leer datos desde SQLite con `pd.read_sql_query()` y desde CSV con `pd.read_csv()`.',
        'Explorar un DataFrame con `head()`, `info()`, `describe()` y `value_counts()` e identificar problemas de calidad.',
        'Construir e interpretar gráficos de distribución, correlación y valores nulos.'],
       ['¿Qué ventaja tiene extraer datos con una consulta SQL dentro de `pd.read_sql_query()` en lugar de leer todas las tablas por separado?',
        '`describe()` muestra una media de S/ 312.57 y una mediana de S/ 199.00. ¿Qué sugiere esa diferencia sobre la distribución?',
        'Una columna numérica aparece como `object` en `info()`. ¿Qué podría estar ocurriendo y por qué es un problema?',
        'La correlación entre visitas y gasto es 0.55. ¿Permite afirmar que aumentar las visitas aumentará el gasto? Justifica.'],
       ['**Semana 5 →** ya sabemos cargar los datos en pandas y explorarlos; el EDA detectó nulos, textos inconsistentes, tipos incorrectos y montos sospechosos.',
        '**Semana 6 →** ahora **corregiremos** esos problemas de forma justificada y guardaremos un conjunto de datos limpio en la base de datos, '
        'listo para construir modelos.'])

# =====================================================================
# SEMANA 6
# =====================================================================
D.h1('Semana 6 · Limpieza de Datos',
     subtitle='Pregunta guía: ¿cómo corrijo los problemas de los datos sin distorsionar la realidad que representan?')

D.h2('6.1 ¿Por qué limpiar?')
D.p('En la Semana 1 vimos el principio *Garbage In, Garbage Out*. La limpieza es la etapa del ciclo de vida del dato en la que se '
    'detectan y tratan los problemas de calidad antes del análisis (Ilyas y Chu, 2019). No consiste en "borrar lo que molesta": cada '
    'decisión cambia los resultados y debe justificarse y documentarse.', align='j')
D.table(['Problema en RutaMarket', 'Dimensión de calidad (Unidad 1)', 'Técnica en pandas'],
        [['Edades vacías, "N.D." o 150', 'Completitud, validez', '`pd.to_numeric(errors=\'coerce\')`, reglas de rango, imputación'],
         ['"Lima", "lima", " LIMA "', 'Consistencia', '`.str.strip().str.title()`'],
         ['Clientes y pedidos registrados dos veces', 'Unicidad', '`duplicated()`, `drop_duplicates()`'],
         ['Métodos de pago faltantes', 'Completitud', '`isnull()`, `fillna()`'],
         ['Precios cien veces mayores que el catálogo', 'Exactitud', 'Detección de atípicos y comparación con tabla de referencia']],
        [5.4, 4.4, 7.2], size=8.5, first_col_bold=True)

D.h2('6.2 Diagnóstico inicial')
py_case('¿Cuántos valores faltan en cada columna?', 's6_nulos',
        '`isnull()` detecta 12 ciudades y 90 métodos de pago faltantes, pero **ninguna edad**. Como vimos en el EDA, las edades vacías están '
        'guardadas como texto (`""` o `"N.D."`) y pandas no las reconoce como nulos. Un diagnóstico basado solo en `isnull()` habría pasado por alto el problema.')

D.h2('6.3 Tipos incorrectos y valores imposibles')
py_case('¿Cuántas edades no son números válidos?', ['s6_tipos', 's6_imposibles'],
        ['`pd.to_numeric(..., errors=\'coerce\')` convierte a número lo que puede y transforma en `NaN` lo que no: aparecen 34 edades faltantes '
         'que antes estaban ocultas. `pd.to_datetime()` convierte `fecha_registro` en fecha, lo que permite calcular antigüedades.',
         'Además, el máximo de 150 años es imposible. Una **regla de validez** (edad entre 18 y 100) lo convierte en faltante: es preferible '
         'reconocer que no se conoce el dato a mantener un valor que se sabe falso.'], index_name='estadístico')

D.h2('6.4 Inconsistencias de texto')
py_case('¿Cuántos clientes hay realmente en cada ciudad?', 's6_texto',
        '`strip()` elimina espacios sobrantes y `title()` unifica el formato ("LIMA" y "lima" pasan a "Lima"). Lima tenía 153 clientes con la '
        'escritura correcta; tras la corrección son 158. Las variantes desaparecen y los conteos por ciudad ya son confiables.', index_name='ciudad')

D.h2('6.5 Registros duplicados')
py_case('¿Hay clientes o pedidos registrados más de una vez?', ['s6_dup_cli', 's6_dup_ped'],
        ['Ocho clientes aparecen dos veces con distinto `id_cliente` pero el mismo correo: la migración les asignó un código nuevo. Por eso los '
         'duplicados se buscan con un criterio de negocio (`subset`), no por el identificador.',
         'En pedidos, 15 registros coinciden en cliente, fecha, monto y estado: corresponden a pedidos registrados dos veces. Antes de eliminar, '
         'conviene confirmar con el área responsable, porque un cliente podría hacer legítimamente dos compras iguales el mismo día.'])

D.h2('6.6 Datos faltantes: detección e imputación')
D.lp('Concepto', '**Imputar** es reemplazar un valor faltante por una estimación razonable. La estrategia depende del tipo de variable y de su distribución:')
D.table(['Estrategia', 'Cuándo usarla', 'Riesgo'],
        [['Eliminar filas', 'Pocos faltantes y sin un patrón', 'Perder información y sesgar la muestra si los faltantes no son al azar'],
         ['Media', 'Variable numérica con distribución simétrica', 'Se desplaza con valores extremos'],
         ['Mediana', 'Variable numérica asimétrica o con atípicos', 'Reduce la variabilidad si hay muchos faltantes'],
         ['Moda', 'Variable categórica', 'Sobrerrepresenta la categoría más frecuente'],
         ['Categoría explícita ("No registrada")', 'Categórica cuando el faltante tiene significado propio', 'Agrega una categoría que no es un valor real']],
        [4.5, 6.0, 6.5], size=8.5, first_col_bold=True)
py_case('¿Qué valores usamos para completar los datos faltantes?', 's6_imputar',
        ['La media (35.6) y la mediana (35.0) de la edad son casi iguales porque su distribución es simétrica; se elige la mediana por ser más '
         'robusta. Para `ciudad` se usa la categoría "No registrada" en lugar de la moda: asignar "Lima" a 12 clientes inventaría ventas en '
         'Lima. Para el método de pago se usa la moda ("Tarjeta"), porque en esta etapa solo se necesita una variable completa.',
         'Cada imputación es un supuesto; debe registrarse para que quien use los datos sepa qué valores son estimados.'])

D.h2('6.7 Valores atípicos: boxplot, z-score e IQR')
D.lp('Concepto', 'Un **valor atípico** (*outlier*) se aleja notablemente del resto. Puede ser un error o un caso real poco frecuente; '
     'detectarlo no equivale a eliminarlo. Hay tres herramientas habituales:')
D.table(['Método', 'Regla', 'Característica'],
        [['Boxplot', 'La caja va de Q1 a Q3; los puntos fuera de los "bigotes" son atípicos', 'Visual; usa la misma regla que el IQR'],
         ['Rango intercuartílico (IQR)', 'IQR = Q3 − Q1; atípico si < Q1 − 1.5·IQR o > Q3 + 1.5·IQR (Tukey, 1977)', 'Robusto: no depende de la media'],
         ['z-score', 'z = (x − media) / desviación estándar; atípico si |z| > 3', 'Supone una distribución aproximadamente simétrica']],
        [3.8, 8.0, 5.2], size=8.5, first_col_bold=True)
py_case('¿Qué pedidos tienen montos atípicos?', ['s6_iqr', 's6_top'],
        ['Los métodos discrepan: el IQR marca 257 pedidos (todo lo que supera S/ 699) y el z-score solo 15. La distribución de montos es '
         'asimétrica: muchos pedidos pequeños y algunos grandes. Por eso el IQR señala como atípicas muchas compras normales de productos '
         'caros, mientras que el z-score, inflado por los valores extremos, solo detecta los más alejados.',
         'Los primeros montos de la lista (S/ 19 070, S/ 12 059…) merecen revisión. Para saber si son errores hay que compararlos con una fuente de referencia.'],
        index=False)
py_case('¿Los montos extremos se deben a precios mal registrados?', ['s6_precio', 's6_corregir'],
        ['Al comparar cada línea con el catálogo aparecen exactamente 6 precios cien veces mayores que el de lista (por ejemplo, 18 990 en lugar '
         'de 189.90): son errores de digitación en los que se omitió el punto decimal. Se corrigen con el precio del catálogo y se recalcula el monto de cada pedido.',
         'Tras la corrección, el máximo baja a S/ 3 178, que corresponde a compras reales de productos caros (por ejemplo, bicicletas estáticas). '
         'Esos pedidos se **conservan**: eliminarlos borraría a los mejores clientes. La Figura 2 muestra el antes y el después.'], index=False)
D.image('fig_boxplot.png', 17.0, 'Figura 2. Boxplot del monto por pedido antes y después de corregir los errores de precio.')

D.h2('6.8 Normalización: Min-Max y estandarización')
D.lp('Concepto', 'Las variables suelen estar en escalas distintas: el número de pedidos va de 1 a 23 y el ticket promedio de decenas a miles '
     'de soles. Los métodos basados en distancias, como K-Means (Semana 8), darían más peso a la variable de mayor escala. **Normalizar** las lleva a una escala común.')
D.table(['', 'Min-Max', 'Estandarización (z-score)'],
        [['Fórmula', "x' = (x − mín) / (máx − mín)", "x' = (x − media) / desviación estándar"],
         ['Resultado', 'Valores entre 0 y 1', 'Media 0 y desviación estándar 1'],
         ['Sensibilidad a atípicos', 'Alta: un extremo comprime al resto', 'También se ve afectada, en menor grado'],
         ['En scikit-learn', '`MinMaxScaler()`', '`StandardScaler()`']],
        [4.0, 6.2, 6.8], size=8.5, first_col_bold=True)
D.p('Para aplicarla necesitamos una tabla con **una fila por cliente**. Se construye igual que un `GROUP BY` de la Unidad 1: contando pedidos, '
    'sumando montos y calculando cuántos días pasaron desde la última compra. Solo se consideran los pedidos entregados.', align='j', keep=True)
py_case(None, ['s6_features', 's6_escalar'],
        'La tabla `modelo` tiene 398 clientes: dos clientes solo tienen pedidos devueltos y no generan gasto. Tras escalar, Min-Max deja ambas '
        'variables entre 0 y 1, y la estandarización, con media 0 y desviación 1. Ninguna variable domina ya por su escala.')

D.h2('6.9 Guardar el dataset limpio en SQLite')
py_case('¿Cómo dejamos los datos limpios disponibles para el resto del equipo?', 's6_guardar',
        ['`to_sql()` guarda cada DataFrame como una tabla nueva en la misma base de datos, sin modificar las tablas originales. Así se conserva '
         'la trazabilidad: los datos originales y los limpios pueden compararse. La verificación con SQL confirma 398 clientes con un gasto medio de S/ 2 112.60.'],
        index=False)
D.table(['Paso', 'Antes', 'Después'],
        [['Clientes', '408 (8 duplicados)', '400'], ['Pedidos', '3 004 (15 duplicados)', '2 989'],
         ['Edades faltantes', '0 detectadas (36 ocultas o imposibles)', '0 (imputadas con la mediana)'],
         ['Ciudades', '18 formas de escribir 6 ciudades y 12 vacías', '6 ciudades + "No registrada"'],
         ['Monto máximo', 'S/ 19 070 (error de digitación)', 'S/ 3 178 (compra real)']],
        [3.6, 6.7, 6.7], size=8.5, first_col_bold=True)

cierre(6,
       ['Diagnosticar nulos, incluidos los "ocultos" como texto, y corregir tipos incorrectos con `to_numeric()` y `to_datetime()`.',
        'Unificar textos inconsistentes y eliminar duplicados con un criterio de negocio.',
        'Elegir entre eliminar, imputar con media, mediana o moda, o usar una categoría explícita, justificando la decisión.',
        'Detectar atípicos con boxplot, IQR y z-score, y distinguir un error de un caso real antes de actuar.',
        'Normalizar variables con Min-Max o estandarización y guardar el resultado en SQLite con `to_sql()`.'],
       ['¿Por qué `isnull()` no detectó las edades vacías? ¿Qué paso previo lo permitió?',
        'Una variable de ingresos tiene media S/ 4 800 y mediana S/ 2 900. ¿Imputarías con la media o con la mediana? ¿Por qué?',
        'El IQR marcó 257 pedidos atípicos. ¿Deberían eliminarse? Explica qué información adicional usarías para decidir.',
        '¿Por qué conviene guardar el dataset limpio como una tabla nueva en lugar de sobrescribir la tabla original?'],
       ['**Semana 6 →** ya tenemos una tabla limpia, `clientes_modelo`, con una fila por cliente, guardada en la base de datos.',
        '**Semana 7 →** con datos confiables podemos construir el primer **modelo de Machine Learning**: estimar el gasto anual de un cliente '
        'a partir de su actividad en la web y evaluar qué tan buena es esa estimación.'])

# =====================================================================
# SEMANA 7
# =====================================================================
D.h1('Semana 7 · ML Supervisado: Regresión y Clasificación',
     subtitle='Pregunta guía: ¿puedo estimar un valor o una categoría a partir de otras variables, y qué tan confiable es esa estimación?')

D.h2('7.1 ¿Qué es Machine Learning?')
D.lp('Concepto', 'El **aprendizaje automático** (*Machine Learning*, ML) agrupa métodos que **aprenden patrones a partir de datos** en lugar de '
     'seguir reglas escritas a mano. En lugar de programar "si visita más de 10 veces, gastará mucho", se entrega al algoritmo un historial '
     'de clientes y él estima la relación (Kelleher et al., 2020).')
D.table(['', 'Aprendizaje supervisado', 'Aprendizaje no supervisado'],
        [['Datos', 'Incluyen la respuesta que se quiere predecir (variable objetivo)', 'No hay variable objetivo'],
         ['Objetivo', 'Predecir un valor o una categoría', 'Descubrir estructura: grupos, patrones'],
         ['Ejemplos de tareas', 'Regresión (valor numérico), clasificación (categoría)', 'Clustering (agrupamiento)'],
         ['En RutaMarket', 'Estimar el gasto anual; identificar clientes de alto valor', 'Segmentar clientes por comportamiento (Semana 8)']],
        [3.2, 7.0, 6.8], size=8.5, first_col_bold=True)

D.h2('7.2 El proceso CRISP-DM')
D.p('Un modelo no empieza con el algoritmo. **CRISP-DM** (*Cross-Industry Standard Process for Data Mining*) es un proceso de referencia '
    'que organiza un proyecto de datos en seis fases iterativas (Chapman et al., 2000). La Unidad 2 lo recorre completo y el proyecto PMD1 lo aplica a un caso real:', align='j', keep=True)
D.image('fig_crisp.png', 17.0, 'Figura 3. Fases de CRISP-DM y su correspondencia con la Unidad 2.')

D.h2('7.3 Datos de entrenamiento y de prueba')
D.lp('Concepto', 'Un modelo se evalúa con datos que **no vio** durante el aprendizaje. Por eso se dividen los datos en un conjunto de '
     '**entrenamiento** (para aprender) y otro de **prueba** (para evaluar), normalmente 80 % y 20 %. Evaluar con los mismos datos del '
     'entrenamiento es como calificar un examen con las preguntas que el estudiante ya conocía.')
py_case('¿Qué datos usaremos para el modelo?', ['s7_leer', 's7_split'],
        'Se leen los datos limpios desde SQLite (Semana 6). La variable objetivo **y** es `gasto_2025`, y la variable explicativa **X**, las visitas '
        'mensuales a la web. `random_state=42` fija la división aleatoria para que el resultado sea reproducible: 318 clientes para entrenar y 80 para evaluar.',
        index_name='estadístico')

D.h2('7.4 Regresión lineal simple')
D.lp('Concepto', 'La **regresión lineal** estima una variable numérica como una recta: **ŷ = b₀ + b₁·x**. El **intercepto** b₀ es el valor '
     'estimado cuando x = 0 y el **coeficiente** b₁ indica cuánto cambia ŷ, en promedio, por cada unidad adicional de x.')
py_case('¿Cuánto gasta al año un cliente según cuántas veces visita la web?', 's7_simple',
        ['**Coeficiente:** cada visita mensual adicional se asocia, en promedio, con **S/ 233.74 más** de gasto anual. Un cliente con 10 visitas '
         'tendría un gasto estimado de 210.05 + 233.74 × 10 ≈ S/ 2 547.',
         '**R² = 0.462:** las visitas explican alrededor del 46 % de las diferencias de gasto entre clientes; el resto depende de factores '
         'que el modelo no incluye. **RMSE = S/ 1 447:** en promedio, la estimación se equivoca en ese monto, grande frente a un gasto medio de '
         'S/ 2 113. El modelo sirve para identificar tendencias, pero no para predecir con precisión el gasto de un cliente individual.'])
D.table(['Métrica', 'Qué mide', 'Cómo leerla'],
        [['R² (coef. de determinación)', 'Proporción de la variación de y explicada por el modelo', 'Entre 0 y 1 (en prueba puede ser negativo); más alto es mejor'],
         ['MSE (error cuadrático medio)', 'Promedio de los errores al cuadrado', 'En unidades al cuadrado (soles²): útil para comparar modelos'],
         ['RMSE (raíz del MSE)', 'Tamaño típico del error', 'En las mismas unidades de y (soles): fácil de comunicar']],
        [4.4, 6.0, 6.6], size=8.5, first_col_bold=True)
D.image('fig_regresion.png', 11.0, 'Figura 4. Gasto anual de cada cliente frente a sus visitas mensuales, con la recta estimada.')
D.p('La Figura 4 muestra por qué el R² es moderado: la tendencia es clara, pero para un mismo número de visitas el gasto varía mucho entre clientes.', align='j')

D.h2('7.5 Regresión lineal múltiple')
py_case('¿Mejora la estimación si agregamos antigüedad y edad?', 's7_multiple',
        ['El R² apenas sube de 0.462 a 0.49 y el RMSE baja de S/ 1 447 a S/ 1 410. Las visitas siguen siendo la variable importante; la antigüedad '
         'y la edad aportan poco. Cada coeficiente se interpreta *manteniendo constantes las demás variables*: el de antigüedad (−17.25) es '
         'pequeño y no debe interpretarse como que la antigüedad reduce el gasto.',
         'La lección para el negocio: agregar variables no garantiza un mejor modelo; conviene agregar las que tengan sentido y verificar la mejora con los datos de prueba.'])

D.h2('7.6 Overfitting y underfitting')
D.table(['', 'Underfitting (subajuste)', 'Buen ajuste', 'Overfitting (sobreajuste)'],
        [['Qué ocurre', 'El modelo es demasiado simple y no capta el patrón', 'Capta el patrón general', 'Memoriza el ruido de los datos de entrenamiento'],
         ['Señal', 'Error alto en entrenamiento y en prueba', 'Errores similares y aceptables', 'Muy buen desempeño en entrenamiento y peor en prueba']],
        [2.6, 4.8, 4.2, 5.4], size=8.5, first_col_bold=True)
py_case('¿Qué pasa si hacemos el modelo más complejo?', 's7_overfit',
        'Con una muestra pequeña de 40 clientes se ajustan curvas cada vez más flexibles (polinomios de grado 1, 3 y 10). Al aumentar la complejidad, '
        'el R² de entrenamiento **sube** (0.54 → 0.69), pero el de prueba **baja** (0.71 → 0.52): el modelo de grado 10 se adapta a las '
        'particularidades de los 28 clientes de entrenamiento y generaliza peor. Este es el sobreajuste, y por eso la evaluación siempre se hace con datos de prueba (James et al., 2021).')

D.h2('7.7 Clasificación: concepto e interpretación')
D.lp('Concepto', 'Cuando la variable objetivo es una **categoría** (sí/no, alto/bajo), el problema es de **clasificación**. El modelo asigna '
     'cada caso a una clase. Aquí definimos como **cliente de alto valor** al 25 % que más gastó y usamos una **regresión logística**, un '
     'clasificador básico, con las mismas tres variables.')
py_case('¿Podemos identificar a los clientes de alto valor por su actividad?', 's7_clasif',
        ['La **exactitud** (accuracy) es 82.5 %: 66 de 80 clientes de prueba fueron clasificados correctamente. Pero la **matriz de confusión** '
         'cuenta otra historia: de los 20 clientes que realmente son de alto valor (fila inferior), el modelo solo reconoce 8 y confunde 12 '
         'con clientes comunes. Cuando el modelo dice "alto valor", acierta en 8 de 10 casos (**precisión** del 80 %), pero solo encuentra '
         'al 40 % de los clientes de alto valor (**sensibilidad** o *recall*).',
         'Si la empresa quiere invitar a todos sus mejores clientes a un programa exclusivo, este modelo dejaría fuera a la mayoría. Con clases '
         'desbalanceadas, una exactitud alta puede ser engañosa: siempre hay que revisar la matriz de confusión.'])
D.table(['', 'Predicho: común (0)', 'Predicho: alto valor (1)'],
        [['Real: común (0)', '58 · verdaderos negativos', '2 · falsos positivos'],
         ['Real: alto valor (1)', '12 · falsos negativos', '8 · verdaderos positivos']],
        [4.4, 6.3, 6.3], size=8.5, first_col_bold=True)

D.h2('7.8 Uso responsable de los modelos')
D.bullets(['**Asociación no es causalidad:** que las visitas se asocien con el gasto no prueba que aumentar las visitas aumente el gasto.',
           '**Calidad de los datos:** un modelo entrenado con los precios mal registrados de la Semana 6 habría aprendido patrones falsos.',
           '**Alcance:** el modelo describe a los clientes de 2025 de RutaMarket; aplicarlo a otro periodo o mercado exige volver a evaluarlo.',
           '**Comunicación:** informar siempre las métricas con su significado para el negocio y las limitaciones del modelo.'])

cierre(7,
       ['Distinguir aprendizaje supervisado y no supervisado, y regresión de clasificación.',
        'Ubicar cada paso de un proyecto de datos en las fases de CRISP-DM.',
        'Dividir los datos en entrenamiento y prueba y explicar por qué se evalúa con datos no vistos.',
        'Entrenar una regresión lineal e interpretar coeficientes, R², MSE y RMSE en términos del negocio.',
        'Reconocer el sobreajuste y leer una matriz de confusión más allá de la exactitud.'],
       ['¿Por qué no es correcto evaluar un modelo con los mismos datos con los que se entrenó?',
        'Un modelo tiene R² = 0.95 en entrenamiento y 0.30 en prueba. ¿Qué problema presenta y qué harías?',
        'Explica a un gerente, sin tecnicismos, qué significa un RMSE de S/ 1 447 en el modelo de gasto.',
        'Un clasificador de fraude tiene 99 % de exactitud, pero solo el 1 % de las transacciones son fraudes. ¿Por qué ese resultado podría no servir?'],
       ['**Semana 7 →** ya sabemos predecir un valor o una categoría cuando conocemos la respuesta en los datos históricos.',
        '**Semana 8 →** ¿y si no hay una respuesta que predecir? Aprenderemos a **descubrir grupos** de clientes con aprendizaje no supervisado (K-Means).'])

# =====================================================================
# SEMANA 8
# =====================================================================
D.h1('Semana 8 · ML No Supervisado: Clustering con K-Means',
     subtitle='Pregunta guía: ¿qué grupos naturales de clientes existen y qué decisión de negocio corresponde a cada uno?')

D.h2('8.1 Aprendizaje no supervisado: ¿cuándo usar clustering?')
D.lp('Concepto', 'El **clustering** (agrupamiento) reúne casos parecidos entre sí y diferentes de los de otros grupos, **sin** una variable objetivo. '
     'El algoritmo no sabe qué grupos existen: los descubre a partir de las variables que se le entregan.')
D.lp('Contexto', 'Se usa cuando se necesita organizar a muchos casos para tratarlos de forma diferenciada: segmentar clientes para campañas, '
     'agrupar productos con patrones de venta similares o identificar tiendas con comportamiento parecido.')
D.lp('Aplicación', 'RutaMarket no puede diseñar una campaña distinta para cada uno de sus 398 clientes, pero sí para cuatro o cinco grupos. '
     'Usaremos tres variables de comportamiento: **cantidad de pedidos**, **ticket promedio** y **días sin comprar**, las mismas que '
     'suelen emplearse en los análisis de recencia, frecuencia y valor.')

D.h2('8.2 Cómo funciona K-Means')
D.p('**K-Means** divide los datos en **k** grupos, cada uno representado por su **centroide** (el punto promedio del grupo) (MacQueen, 1967):', align='j', keep=True)
D.numbered(['Se eligen k centroides iniciales al azar.',
            'Cada cliente se asigna al centroide más cercano.',
            'Cada centroide se recalcula como el promedio de los clientes asignados.',
            'Se repiten los pasos 2 y 3 (**iteraciones**) hasta que las asignaciones dejan de cambiar: el algoritmo **converge**.'])
D.p('La **inercia** es la suma de las distancias al cuadrado de cada cliente a su centroide: mide qué tan compactos son los grupos. Como el '
    'resultado depende del punto de partida, `n_init=10` ejecuta el algoritmo diez veces y conserva la mejor solución.', align='j')

D.h2('8.3 Preparar los datos y elegir k: el método del codo')
D.p('K-Means usa distancias, por lo que las variables se **estandarizan** primero (Semana 6). Para elegir k se calcula la inercia con distintos '
    'valores: siempre disminuye al agregar grupos, pero a partir de cierto punto la mejora es pequeña. Ese punto de quiebre es el **codo**.', align='j', keep=True)
py_case('¿Cuántos segmentos conviene formar?', 's8_escalar',
        'La inercia cae con fuerza de k = 1 a k = 4 (de 1 194 a 322) y luego las reducciones son menores (258, 224…). Elegimos **k = 4**: '
        'agrega poca complejidad y los grupos siguen siendo interpretables. El codo es una guía, no una regla exacta; la decisión final '
        'también considera si los grupos tienen sentido para el negocio.')

D.h2('8.4 Interpretar los clusters')
py_case('¿Cómo es cada grupo de clientes?', ['s8_kmeans', 's8_nombrar'], None, index_name='')
D.image('fig_kmeans.png', 17.0, 'Figura 5. Método del codo (izquierda) y clientes según pedidos y días sin comprar, por segmento (derecha).')
D.p('El algoritmo solo entrega números de cluster; **nombrarlos e interpretarlos** es tarea del analista, a partir de los promedios de cada grupo:', align='j', keep=True)
D.table(['Segmento', 'Perfil (promedios)', 'Decisión de negocio sugerida'],
        [['Frecuentes (171)', '12.2 pedidos, ticket S/ 255, última compra hace 30 días', 'Programa de fidelización y beneficios por recurrencia'],
         ['Regulares (114)', '4.7 pedidos, ticket S/ 401, última compra hace 61 días', 'Recomendaciones personalizadas para aumentar la frecuencia'],
         ['Inactivos (98)', '2.1 pedidos, ticket S/ 208, sin comprar hace 268 días', 'Campaña de reactivación con incentivo acotado'],
         ['Compras grandes ocasionales (15)', '1.8 pedidos, ticket S/ 1 029, sin comprar hace 212 días', 'Atención personalizada y avisos de productos de alto valor']],
        [4.2, 6.4, 6.4], size=8.5, first_col_bold=True)

D.h2('8.5 Guardar los segmentos en la base de datos')
py_case('¿Cuánto aporta cada segmento a las ventas?', 's8_guardar',
        ['El segmento queda guardado como una **nueva columna** de `clientes_modelo`, de modo que cualquier consulta SQL posterior puede usarlo, '
         'como la verificación mostrada. Los 171 clientes frecuentes generan S/ 551 530, el 65.6 % del gasto total; los 98 inactivos solo el 5.1 %.',
         'Para la gerencia, el mensaje es claro: proteger a los clientes frecuentes es prioritario y reactivar a los inactivos tiene potencial, '
         'pero partiendo de un gasto bajo. Así, el resultado del análisis vuelve a la base de datos y queda disponible para otras áreas.'], index=False)

D.h2('8.6 Limitaciones de K-Means')
D.bullets(['Hay que elegir k de antemano; el método del codo ayuda, pero no siempre muestra un quiebre claro.',
           'Es sensible a la escala (por eso se estandariza) y a los valores atípicos, que pueden desplazar los centroides.',
           'Tiende a formar grupos compactos y de forma aproximadamente esférica; grupos alargados o de tamaños muy distintos pueden quedar mal separados.',
           'Los segmentos describen comportamientos pasados: deben actualizarse cuando cambian los datos.'])

cierre(8,
       ['Explicar cuándo usar clustering y en qué se diferencia de un modelo supervisado.',
        'Describir el funcionamiento de K-Means: centroides, asignación, iteraciones y convergencia.',
        'Estandarizar variables antes de aplicar K-Means y justificar la elección de k con el método del codo.',
        'Interpretar los perfiles de cada cluster, nombrarlos y proponer decisiones de negocio.',
        'Guardar los segmentos en la base de datos y reconocer las limitaciones del método.'],
       ['¿Por qué K-Means necesita que las variables estén en la misma escala? ¿Qué ocurriría con el ticket promedio si no se estandariza?',
        'En el método del codo, ¿por qué no se elige simplemente el k con menor inercia?',
        'Un cluster tiene ticket alto pero pocos pedidos y mucho tiempo sin comprar. ¿Qué acción comercial propondrías y cómo medirías su resultado?',
        '¿Qué ventaja tiene guardar el segmento como columna en la base de datos en lugar de dejarlo solo en el notebook?'],
       ['**Semana 8** es también la semana de evaluación de la unidad (Pa2) y de entrega del proyecto PMD1, que integra todo lo aprendido. '
        'La siguiente sección resume ese recorrido completo.'], siguiente='Hacia la evaluación de la Unidad 2')

# =====================================================================
# INTEGRACIÓN
# =====================================================================
D.h1('Integración de la Unidad 2', subtitle='De la base de datos a una decisión respaldada por datos limpios y modelos interpretables.')
D.p('Las cuatro semanas forman un solo flujo de trabajo, que continúa el de la Unidad 1. Cada etapa depende de la anterior: un modelo '
    'entrenado con datos sucios produce conclusiones falsas, y un modelo sin interpretación no ayuda a decidir.', align='j')
D.h2('Mapa del recorrido')
D.table(['Etapa', 'Herramienta', 'Semana', 'Pregunta que responde'],
        [['Extraer', 'SQL + `pd.read_sql_query()`', 'U1 · S5', '¿Qué datos necesito y cómo los combino?'],
         ['→ Explorar (EDA)', '`info()`, `describe()`, `value_counts()`, gráficos', 'S5', '¿Cómo son los datos y qué problemas tienen?'],
         ['→ Limpiar', '`to_numeric()`, `drop_duplicates()`, `fillna()`, IQR', 'S6', '¿Cómo corrijo los problemas sin distorsionar la realidad?'],
         ['→ Transformar', '`groupby()`, `StandardScaler()`', 'S6', '¿Qué tabla y qué escala necesita el modelo?'],
         ['→ Predecir', '`LinearRegression()`, `LogisticRegression()`', 'S7', '¿Puedo estimar un valor o una categoría?'],
         ['→ Segmentar', '`KMeans()`', 'S8', '¿Qué grupos existen?'],
         ['→ Guardar y decidir', '`to_sql()` + SQL', 'S6 · S8', '¿Cómo pongo el resultado al servicio de la organización?']],
        [3.3, 6.2, 1.8, 5.7], size=8.4, first_col_bold=True, align=['l', 'l', 'c', 'l'])

D.h2('Caso integrador')
D.p('La gerencia de RutaMarket dispone de un presupuesto limitado para una **campaña de reactivación** y pregunta: *¿a quiénes debemos '
    'dirigirla y qué podemos esperar?* Responderla requiere todo el recorrido de la unidad:', align='j')
D.lp('1. Explorar (Semana 5)', 'El EDA muestra que la base tiene clientes con muy distinta actividad y detecta problemas: duplicados, edades '
     'ocultas como texto y montos sospechosos. Sin corregirlos, contaríamos clientes dos veces y sobrestimaríamos el gasto.')
D.lp('2. Limpiar (Semana 6)', 'Tras eliminar duplicados y corregir los precios mal digitados, el gasto de 2025 es confiable y se construye la '
     'tabla `clientes_modelo`, con días sin comprar, pedidos y ticket promedio.')
D.lp('3. Segmentar (Semana 8)', 'K-Means identifica 98 clientes **inactivos** (sin comprar hace unos 268 días en promedio) y 15 clientes de '
     '**compras grandes ocasionales** (ticket de S/ 1 029). Son los candidatos naturales a la campaña, con mensajes distintos para cada grupo.')
D.lp('4. Estimar (Semana 7)', 'El modelo de regresión indica que el gasto se asocia con las visitas a la web. Por eso, un indicador razonable '
     'del éxito de la campaña es el aumento de las visitas de los clientes reactivados, sin prometer un monto exacto: el RMSE recuerda que las '
     'estimaciones individuales son imprecisas.')
D.lp('5. Decidir y comunicar', 'La recomendación se sustenta en datos: priorizar a los inactivos con un incentivo acotado y ofrecer atención '
     'personalizada a los compradores de alto ticket, midiendo el resultado con los mismos indicadores en el siguiente periodo. El segmento queda '
     'guardado en la base de datos para que el área de marketing lo consulte.')

D.h2('Relación con el proyecto PMD1')
D.table(['Fase del proyecto', 'Semana', 'Qué aplica de esta unidad'],
        [['Fase 1: problema de negocio + EDA', '6 (checkpoint)', 'Carga del dataset en SQLite, lectura con pandas, EDA e interpretación de gráficos (Semana 5)'],
         ['Fase 2: limpieza y modelado', '7', 'Pipeline de limpieza documentado (Semana 6) y modelo supervisado evaluado con datos de prueba (Semana 7)'],
         ['Fase 3: entrega final', '8', 'Comparación de modelos, segmentación (Semana 8), informe gerencial con recomendaciones y código reproducible']],
        [4.8, 2.4, 9.8], size=8.5, first_col_bold=True, align=['l', 'c', 'l'])

D.h2('Checklist de dominio de la Unidad 2')
D.p('Marca cada afirmación solo si puedes demostrarla en un notebook propio:', keep=True)
D.checklist(['Leo datos desde SQLite con una consulta SQL y desde un archivo CSV, y los cargo en un DataFrame.',
             'Exploro un DataFrame con `info()`, `describe()` y `value_counts()` y redacto los hallazgos en celdas de texto.',
             'Construyo e interpreto gráficos de distribución, correlación y valores nulos.',
             'Detecto nulos ocultos, tipos incorrectos, textos inconsistentes y duplicados, y los corrijo justificando cada decisión.',
             'Elijo una estrategia de imputación (media, mediana, moda o categoría explícita) según el tipo y la distribución de la variable.',
             'Detecto atípicos con boxplot, IQR y z-score, y verifico si son errores antes de modificarlos.',
             'Normalizo variables y guardo los datos limpios como nuevas tablas en SQLite.',
             'Entreno una regresión lineal con datos de entrenamiento y prueba e interpreto coeficientes, R², MSE y RMSE.',
             'Reconozco el sobreajuste y leo una matriz de confusión más allá de la exactitud.',
             'Aplico K-Means, elijo k con el método del codo, interpreto los segmentos y los guardo en la base de datos.'])

D.h1('Conclusiones', page_break=False)
D.numbered(['Python y pandas no reemplazan a SQL, lo complementan: SQL extrae y combina los datos en la base y pandas los explora, '
            'transforma y prepara para el análisis. El flujo termina devolviendo los resultados a la base de datos.',
            'El análisis exploratorio es el primer control de calidad: permitió descubrir en RutaMarket nulos ocultos como texto, '
            'ciudades escritas de varias formas, duplicados y montos imposibles antes de que distorsionaran cualquier conclusión.',
            'Limpiar datos exige criterio: un valor atípico puede ser un error o un buen cliente, y cada imputación es un supuesto. '
            'Las decisiones deben justificarse con el negocio y documentarse.',
            'Un modelo supervisado se evalúa con datos que no vio y se interpreta en el lenguaje del negocio. Métricas como R², RMSE o la '
            'matriz de confusión muestran no solo cuánto acierta el modelo, sino también en qué se equivoca.',
            'El clustering convierte datos de comportamiento en segmentos accionables, pero sus grupos solo tienen valor cuando el analista '
            'los interpreta, les da nombre y los vincula a decisiones concretas.',
            'La calidad de un modelo depende de la calidad de los datos y de su gestión: sin la base construida en la Unidad 1, ninguno de estos análisis sería confiable.'],
           after=5)

D.h1('Referencias bibliográficas', page_break=False)
refs = ['Bruce, P., Bruce, A., y Gedeck, P. (2020). *Practical statistics for data scientists* (2.ª ed.). O\'Reilly Media.',
        'Chapman, P., Clinton, J., Kerber, R., Khabaza, T., Reinartz, T., Shearer, C., y Wirth, R. (2000). *CRISP-DM 1.0: Step-by-step data mining guide*. SPSS.',
        'Gironés Roig, J., Casas Roma, J., Minguillón Alfonso, J., y Caihuelas Quiles, R. (2017). *Minería de datos: modelos y algoritmos*. Editorial UOC.',
        'Grus, J. (2019). *Data science from scratch: First principles with Python* (2.ª ed.). O\'Reilly Media.',
        'Hunter, J. D. (2007). Matplotlib: A 2D graphics environment. *Computing in Science & Engineering, 9*(3), 90–95.',
        'Ilyas, I. F., y Chu, X. (2019). *Data cleaning*. ACM Books.',
        'James, G., Witten, D., Hastie, T., y Tibshirani, R. (2021). *An introduction to statistical learning* (2.ª ed.). Springer.',
        'Kelleher, J. D., Mac Namee, B., y D\'Arcy, A. (2020). *Fundamentals of machine learning for predictive data analytics* (2.ª ed.). MIT Press.',
        'MacQueen, J. (1967). Some methods for classification and analysis of multivariate observations. En *Proceedings of the Fifth Berkeley Symposium on Mathematical Statistics and Probability* (Vol. 1, pp. 281–297). University of California Press.',
        'McKinney, W. (2022). *Python for data analysis: Data wrangling with pandas, NumPy, and Jupyter* (3.ª ed.). O\'Reilly Media.',
        'Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., … Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. *Journal of Machine Learning Research, 12*, 2825–2830.',
        'Provost, F., y Fawcett, T. (2013). *Data science for business*. O\'Reilly Media.',
        'Tukey, J. W. (1977). *Exploratory data analysis*. Addison-Wesley.',
        'Waskom, M. L. (2021). seaborn: Statistical data visualization. *Journal of Open Source Software, 6*(60), 3021. https://doi.org/10.21105/joss.03021']
for r in refs:
    para = D.p(r, size=9.5, after=4)
    para.paragraph_format.left_indent = Cm(1.0); para.paragraph_format.first_line_indent = Cm(-1.0)

D.footer('Fundamentos de Gestión de Datos · Unidad 2: Python para Datos, Limpieza y ML Básico · TECSUP 2026-II')
D.save(OUT)
print('ok', OUT)
