"""Helpers de formato para construir el documento con python-docx."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
TEAL = RGBColor(0x2E, 0x75, 0xB6)
GRAY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
FONT = 'Arial'
MONO = 'Consolas'
BODY = 10
TEXT_W = 17.0  # ancho útil en cm (A4 21 cm - 2 x 2 cm)


def _shade(el_pr, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    el_pr.append(shd)


def _borders(tc_pr, spec):
    """spec: dict side -> (size_eighths, color) o None para sin borde."""
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}')
        if spec.get(side):
            sz, col = spec[side]
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), col)
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    tc_pr.append(b)


def _cell_margins(tbl, top=60, bottom=60, left=100, right=100):
    tbl_pr = tbl._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tbl_pr.append(m)


def _no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); tr_pr.append(e)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); tr_pr.append(e)


def _keep_next(p):
    p.paragraph_format.keep_with_next = True


INLINE = re.compile(r'(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)')


def add_runs(p, text, size=None, color=None, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith('**') and part.endswith('**'):
            part = part[2:-2]; b = True
        elif part.startswith('`') and part.endswith('`'):
            part = part[1:-1]; mono = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part = part[1:-1]; i = True
        r = p.add_run(part)
        r.bold = b; r.italic = i
        if mono:
            r.font.name = MONO
            r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        if size:
            r.font.size = Pt(size if not mono else size - 0.5)
        elif mono:
            r.font.size = Pt(BODY - 1)
        if color is not None and not mono:
            r.font.color.rgb = color
    return p


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = Cm(2); s.bottom_margin = Cm(1.8)
        s.header_distance = Cm(1); s.footer_distance = Cm(0.9)
        st = self.d.styles
        n = st['Normal']
        n.font.name = FONT; n.font.size = Pt(BODY)
        n.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        n.paragraph_format.space_after = Pt(5)
        n.paragraph_format.line_spacing = 1.08
        for name, size, color, before, after in (('Heading 1', 17, NAVY, 0, 8),
                                                  ('Heading 2', 13, NAVY, 12, 5),
                                                  ('Heading 3', 11, TEAL, 8, 3)):
            h = st[name]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
            h.font.color.rgb = color; h.font.italic = False
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn('w:rFonts'))
            if rf is None:
                rf = OxmlElement('w:rFonts'); rpr.append(rf)
            for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                rf.set(qn(a), FONT)
            for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True

    # ---------- bloques de texto ----------
    def h1(self, text, page_break=True, subtitle=None):
        p = self.d.add_paragraph(style='Heading 1')
        p.paragraph_format.page_break_before = page_break
        add_runs(p, text)
        pf = p.paragraph_format
        # línea inferior
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '12'); bot.set(qn('w:space'), '4'); bot.set(qn('w:color'), '2E75B6')
        bdr.append(bot); ppr.append(bdr)
        if subtitle:
            q = self.d.add_paragraph()
            add_runs(q, subtitle, size=10, color=GRAY, italic=True)
            q.paragraph_format.space_after = Pt(8)
        return p

    def h2(self, text):
        p = self.d.add_paragraph(style='Heading 2'); add_runs(p, text); return p

    def h3(self, text):
        p = self.d.add_paragraph(style='Heading 3'); add_runs(p, text); return p

    def p(self, text, size=None, color=None, italic=False, align=None, after=None, keep=False):
        para = self.d.add_paragraph()
        add_runs(para, text, size=size, color=color, italic=italic)
        if align == 'j':
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == 'c':
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if keep:
            _keep_next(para)
        return para

    def lp(self, label, text, color=TEAL):
        """Párrafo con etiqueta inicial (Concepto, Contexto, ...)."""
        para = self.d.add_paragraph()
        r = para.add_run(label + '. '); r.bold = True; r.font.color.rgb = color
        add_runs(para, text)
        para.paragraph_format.space_after = Pt(4)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def bullets(self, items, size=None, after=2, indent=0.5):
        for it in items:
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(indent + 0.4); pf.first_line_indent = Cm(-0.4)
            pf.space_after = Pt(after)
            r = para.add_run('•\t'); r.font.color.rgb = TEAL; r.bold = True
            pf.tab_stops.add_tab_stop(Cm(indent + 0.4))
            add_runs(para, it, size=size)

    def numbered(self, items, size=None, after=2, indent=0.5):
        for k, it in enumerate(items, 1):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(indent + 0.55); pf.first_line_indent = Cm(-0.55)
            pf.space_after = Pt(after)
            pf.tab_stops.add_tab_stop(Cm(indent + 0.55))
            r = para.add_run(f'{k}.\t'); r.bold = True; r.font.color.rgb = TEAL
            add_runs(para, it, size=size)

    def checklist(self, items):
        for it in items:
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(1.0); pf.first_line_indent = Cm(-0.6); pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Cm(1.0))
            r = para.add_run('☐\t'); r.font.color.rgb = TEAL
            add_runs(para, it)

    def image(self, path, width_cm, caption=None):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(path, width=Cm(width_cm))
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.keep_with_next = bool(caption)
        if caption:
            c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(c, caption, size=8.5, color=GRAY, italic=True)
            c.paragraph_format.space_after = Pt(6)

    def spacer(self, pt=4):
        para = self.d.add_paragraph(); para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = Pt(pt)
        return para

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- tablas ----------
    def table(self, headers, rows, widths, size=8.8, header_fill='1F3864', zebra=True,
              first_col_bold=False, align=None, keep=True):
        t = self.d.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _cell_margins(t, 45, 45, 90, 90)
        total = sum(widths)
        for i, w in enumerate(widths):
            for cell in t.columns[i].cells:
                cell.width = Cm(w)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths[i] * 567)))
        light = 'BFBFBF'
        for ri, row in enumerate(t.rows):
            _no_split(row)
            vals = headers if ri == 0 else rows[ri - 1]
            for ci, cell in enumerate(row.cells):
                tcpr = cell._tc.get_or_add_tcPr()
                if ri == 0:
                    _shade(tcpr, header_fill)
                elif zebra and ri % 2 == 0:
                    _shade(tcpr, 'EEF3FA')
                _borders(tcpr, {'top': (4, light), 'bottom': (4, light), 'left': (4, light), 'right': (4, light)})
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                if keep and ri < len(t.rows) - 1:
                    para.paragraph_format.keep_with_next = True
                txt = '' if vals[ci] is None else str(vals[ci])
                if align and align[ci] == 'r':
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif align and align[ci] == 'c':
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if ri == 0:
                    add_runs(para, txt, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                else:
                    add_runs(para, txt, size=size, bold=(first_col_bold and ci == 0))
            if ri == 0:
                _repeat_header(row)
        self.spacer(5)
        return t

    def box(self, title, lines=(), fill='EEF3FA', edge='2E75B6', bullets=None, size=9.5, title_color=NAVY):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cell_margins(t, 90, 90, 160, 140)
        cell = t.rows[0].cells[0]
        cell.width = Cm(TEXT_W)
        t._tbl.tblGrid.findall(qn('w:gridCol'))[0].set(qn('w:w'), str(int(TEXT_W * 567)))
        _no_split(t.rows[0])
        tcpr = cell._tc.get_or_add_tcPr()
        _shade(tcpr, fill)
        _borders(tcpr, {'left': (24, edge)})
        first = cell.paragraphs[0]
        paras = []
        if title:
            add_runs(first, title, size=size + 0.5, color=title_color, bold=True)
            paras.append(first); first = None
        for ln in lines:
            para = first if first is not None else cell.add_paragraph()
            first = None
            add_runs(para, ln, size=size)
            paras.append(para)
        for it in (bullets or []):
            para = first if first is not None else cell.add_paragraph()
            first = None
            pf = para.paragraph_format
            pf.left_indent = Cm(0.45); pf.first_line_indent = Cm(-0.35)
            pf.tab_stops.add_tab_stop(Cm(0.45))
            r = para.add_run('•\t'); r.font.color.rgb = TEAL; r.bold = True; r.font.size = Pt(size)
            add_runs(para, it, size=size)
            paras.append(para)
        for k, para in enumerate(paras):
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.08
            if k < len(paras) - 1:
                para.paragraph_format.keep_with_next = True
        self.spacer(6)
        return t

    def code(self, text, size=8.6, fill='F4F6F9', edge='7F9CC4'):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cell_margins(t, 70, 70, 160, 120)
        cell = t.rows[0].cells[0]
        cell.width = Cm(TEXT_W)
        t._tbl.tblGrid.findall(qn('w:gridCol'))[0].set(qn('w:w'), str(int(TEXT_W * 567)))
        long = text.count('\n') > 14
        if not long:
            _no_split(t.rows[0])
        tcpr = cell._tc.get_or_add_tcPr()
        _shade(tcpr, fill)
        _borders(tcpr, {'left': (18, edge)})
        lines = text.split('\n')
        for k, ln in enumerate(lines):
            para = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.keep_with_next = not long
            comment = ln.strip().startswith('--')
            r = para.add_run(ln if ln else ' ')
            r.font.name = MONO; r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0x38, 0x76, 0x1D) if comment else RGBColor(0x1A, 0x1A, 0x1A)
        self.spacer(4)
        return t

    def label(self, text, color=TEAL):
        para = self.d.add_paragraph()
        r = para.add_run(text.upper()); r.bold = True; r.font.size = Pt(8.2); r.font.color.rgb = color
        para.paragraph_format.space_after = Pt(1); para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.keep_with_next = True
        # espaciado entre letras
        rpr = r._element.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '10'); rpr.append(sp)
        return para

    # ---------- pie de página ----------
    def footer(self, left_text):
        s = self.d.sections[0]
        s.different_first_page_header_footer = True
        f = s.footer.paragraphs[0]
        f.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), alignment=2)
        r = f.add_run(left_text + '\t'); r.font.size = Pt(8); r.font.color.rgb = GRAY
        r2 = f.add_run(); r2.font.size = Pt(8); r2.font.color.rgb = NAVY; r2.bold = True
        for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
            if kind:
                e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind); r2._element.append(e)
            else:
                e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r2._element.append(e)
        ppr = f._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr'); top = OxmlElement('w:top')
        top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '4'); top.set(qn('w:space'), '4'); top.set(qn('w:color'), 'BFBFBF')
        bdr.append(top); ppr.append(bdr)

    def save(self, path):
        self.d.save(path)
